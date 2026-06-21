#!/usr/bin/env python3
"""
Import the ECOSLO "Tree Planting Log" workbook into Supabase-ready CSV files.

Reads three sheets from an Excel file shaped like and emits CSVs that can be
imported directly into Supabase:

  * trees.csv         -> public.trees
  * members.csv       -> public.members   (adopters / tree keepers)
  * skipped_rows.csv  -> rows that could not be imported, with the reason

Usage:
    python scripts/import_trees.py "Tree Planting Log.xlsx" [-o OUTPUT_DIR]

--------------------------------------------------------------------------------
Mapping / business rules
--------------------------------------------------------------------------------
Sheets imported (matched case-insensitively, punctuation-insensitively):
    "2526 Planting Season"  -> status Active     (layout A)
    "Planted Log"           -> status Active     (layout B)
    "Off-Boarded Trees"     -> status Graduated  (layout B)
Any other sheet (Dead Trees, Pivot Table, ReLeaf Final List, ...) is ignored.

trees.condition <- the sheet "Status" column (the tree's health), lowercased:
    Good->good, Fair->fair, Okay->okay, Decent->decent, Dead->dead, blank->good

is_public <- "Public/Private" column: starts with "Pub" -> true, else false.

trees.notes <- the sheet "Notes" column, with every other unmapped column
    (ReLeaf #, FullCircle #, stock size, neighborhood, census tract, low income,
    planting area, distance/direction, tagged, mulched, immediate needs, ...)
    appended as labeled lines. Blank / "N/A" / "n/a" values are dropped.

Required columns (NOT NULL with no DB default). A row missing any of these is
skipped and written to skipped_rows.csv instead of being imported:
    ecoslo_num, species_name, common_name, funder, date_planted, address,
    latitude, longitude

De-duplication: ecoslo_num is UNIQUE, so only one row is kept per tree number.
    Priority Off-Boarded (3) > 25/26 (2) > Planted Log (1); ties keep the first
    occurrence. De-dup runs over *valid* rows only, so if the higher-priority
    row was skipped (e.g. missing coordinates) a valid lower-priority row is
    still imported (this happens for tree #205). Dropped duplicates are listed
    in skipped_rows.csv.

members.csv: adopter / tree-keeper info from the kept trees, de-duplicated to
    one member per person. Key = email when present, otherwise normalized name.
    Name is split first-word -> firstname, remainder -> lastname. Adopters with
    no email get a unique placeholder ``noemail+N@ecoslo.invalid`` and
    needs_review=true (email is NOT NULL + UNIQUE in public.members).
    trees_assigned is the aggregated set of that member's ecoslo_num values.
    trees.tree_keeper_id is intentionally left unset (members.id is generated on
    import and unknown ahead of time); the link is carried by
    members.trees_assigned, which the app's RLS policies already use.
"""

import argparse
import csv
import os
import re
import sys
import warnings
from datetime import date, datetime

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

# openpyxl warns on a corrupt cached date in the workbook's Census Tract column;
# that column is only folded into notes, so the warning is noise here.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Values that mean "no data" across the workbook.
JUNK = {"", "n/a", "na", "n\\a", "none", "-", "--", "did not provide", "unknown", "?"}

# ---------------------------------------------------------------------------
# Sheet configuration
# ---------------------------------------------------------------------------
# layout A: "2526 Planting Season"      layout B: "Planted Log" / "Off-Boarded"
# Matched against a normalized sheet name (lowercase, alphanumerics only).
SHEET_RULES = [
    # (substring-in-normalized-name, layout, status, priority)
    ("offboard",       "B", "Graduated", 3),
    ("plantedlog",     "B", "Active",    1),
    ("plantingseason", "A", "Active",    2),
]

# Column indices (0-based) per layout.
COLS = {
    "A": {
        "status": 0, "tree_num": 1, "species": 2, "common": 3, "funder": 4,
        "date": 5, "address": 6, "latlong": 7, "pubpriv": 8,
        "adopter_name": 9, "adopter_phone": 10, "adopter_email": 11,
        "notes": 14,
        # extra columns folded into notes: (label, index, kind)
        "extras": [
            ("Watered this week", 12, "text"),
            ("Next mulching due date", 13, "text"),
        ],
    },
    "B": {
        "status": 0, "species": 2, "common": 3, "tree_num": 4, "funder": 7,
        "date": 8, "address": 9, "lat": 10, "long": 11, "pubpriv": 16,
        "adopter_name": 20, "adopter_phone": 21, "adopter_email": 22,
        "notes": 27,
        "extras": [
            ("ReLeaf #", 5, "text"),
            ("FullCircle #", 6, "text"),
            ("Stock size", 12, "text"),
            ("Neighborhood", 13, "text"),
            ("Census tract", 14, "int"),
            ("Low income", 15, "text"),
            ("Planting area", 17, "text"),
            ("Distance to nearest building", 18, "text"),
            ("Direction from building", 19, "text"),
            ("Tagged", 23, "bool"),
            ("Mulched Jan 2023", 24, "bool"),
            ("Immediate needs", 26, "text"),
        ],
    },
}

CONDITIONS = {"good", "fair", "poor", "okay", "decent", "dead"}


# ---------------------------------------------------------------------------
# Value helpers
# ---------------------------------------------------------------------------
def clean(v):
    return "" if v is None else str(v).strip()


def is_junk(v):
    return clean(v).lower() in JUNK


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")


def norm_ws(v):
    """Collapse all whitespace (including newlines) to single spaces."""
    return re.sub(r"\s+", " ", clean(v)).strip()


def extract_email(v):
    """Return the first valid-looking email found in a messy cell, lowercased.

    Adopter email cells often hold more than an address: a second address
    (``a@x.com + b@y.com``), a parenthetical (``brenna (b@x.org)``), or a note
    on a new line. We key members on the first real address found.
    """
    m = EMAIL_RE.search(clean(v))
    return m.group(0).lower() if m else ""


def to_int(v):
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def to_float(v):
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def to_date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = clean(v)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%m-%d-%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_combined_latlong(v):
    """Parse '(35.31, -120.83)' -> (35.31, -120.83)."""
    nums = re.findall(r"-?\d+(?:\.\d+)?", clean(v))
    if len(nums) < 2:
        return (None, None)
    return (to_float(nums[0]), to_float(nums[1]))


def to_condition(v):
    s = clean(v).lower()
    if not s:
        return "good", None
    if s in CONDITIONS:
        return s, None
    return "good", f"unrecognized condition {clean(v)!r} mapped to good"


def to_is_public(v):
    return clean(v).lower().startswith("pub")


def fmt_extra(v, kind):
    if v is None:
        return ""
    if kind == "bool":
        if isinstance(v, bool):
            return "Yes" if v else "No"
        s = clean(v).lower()
        if s in ("true", "yes", "y", "1"):
            return "Yes"
        if s in ("false", "no", "n", "0"):
            return "No"
        return clean(v)
    if kind == "int":
        n = to_int(v)
        return str(n) if n is not None else clean(v)
    return clean(v)


def build_notes(row, cfg):
    base = clean(row[cfg["notes"]])
    if base.lower() in JUNK:
        base = ""
    lines = []
    for label, idx, kind in cfg["extras"]:
        if idx >= len(row):
            continue
        val = fmt_extra(row[idx], kind)
        if val and val.lower() not in JUNK:
            lines.append(f"{label}: {val}")
    if lines:
        block = "\n".join(lines)
        return f"{base}\n\n{block}" if base else block
    return base


# ---------------------------------------------------------------------------
# Core parsing
# ---------------------------------------------------------------------------
def match_sheet(name):
    norm = re.sub(r"[^a-z0-9]", "", name.lower())
    for needle, layout, status, priority in SHEET_RULES:
        if needle in norm:
            return layout, status, priority
    return None


def row_is_empty(row):
    return not any(c is not None and str(c).strip() != "" for c in row)


def parse_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    valid = []      # list of tree record dicts
    skipped = []    # list of skipped-row dicts
    warns = []      # non-fatal notes
    sheet_counts = {}

    for sheet_name in wb.sheetnames:
        match = match_sheet(sheet_name)
        if not match:
            continue
        layout, status, priority = match
        cfg = COLS[layout]
        ws = wb[sheet_name]
        n_valid = n_skip = 0

        for excel_row, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if excel_row == 1 or row_is_empty(row):
                continue
            row = list(row)

            raw_num = row[cfg["tree_num"]] if cfg["tree_num"] < len(row) else None
            species = clean(row[cfg["species"]]) if cfg["species"] < len(row) else ""
            common = clean(row[cfg["common"]]) if cfg["common"] < len(row) else ""
            raw_funder = row[cfg["funder"]] if cfg["funder"] < len(row) else None
            # funder is NOT NULL with no default: blank, junk ("n/a"), or a stray
            # date value (both seen in the workbook) all count as missing.
            funder = "" if isinstance(raw_funder, (datetime, date)) or is_junk(raw_funder) else clean(raw_funder)
            address = clean(row[cfg["address"]]) if cfg["address"] < len(row) else ""
            raw_date = row[cfg["date"]] if cfg["date"] < len(row) else None

            if layout == "A":
                raw_lat = raw_long = row[cfg["latlong"]] if cfg["latlong"] < len(row) else None
                lat, lon = parse_combined_latlong(raw_lat)
            else:
                raw_lat = row[cfg["lat"]] if cfg["lat"] < len(row) else None
                raw_long = row[cfg["long"]] if cfg["long"] < len(row) else None
                lat, lon = to_float(raw_lat), to_float(raw_long)

            num = to_int(raw_num)
            dp = to_date(raw_date)

            reasons = []
            if num is None:
                reasons.append("missing/invalid tree number")
            if is_junk(species):
                reasons.append("missing species")
            if is_junk(common):
                reasons.append("missing common name")
            if not funder:
                reasons.append("missing funder")
            if not dp:
                reasons.append("missing/invalid date planted")
            if is_junk(address):
                reasons.append("missing address")
            if lat is None or lon is None:
                reasons.append("missing/invalid lat/long")
            elif not (-90 <= lat <= 90 and -180 <= lon <= 180):
                reasons.append("lat/long out of range")

            condition, cwarn = to_condition(row[cfg["status"]] if cfg["status"] < len(row) else None)
            if cwarn:
                warns.append(f"{sheet_name} row {excel_row}: {cwarn}")

            skip_entry = {
                "source_sheet": sheet_name,
                "excel_row": excel_row,
                "raw_tree_num": clean(raw_num),
                "species": species,
                "common_name": common,
                "funder": funder,
                "date_planted": clean(raw_date),
                "address": address,
                "raw_lat": clean(raw_lat),
                "raw_long": clean(raw_long),
                "status_intended": status,
            }

            if reasons:
                skip_entry["reason"] = "; ".join(reasons)
                skipped.append(skip_entry)
                n_skip += 1
                continue

            record = {
                "ecoslo_num": num,
                "status": status,
                "species_name": species,
                "common_name": common,
                "funder": funder,
                "date_planted": dp,
                "address": address,
                "latitude": lat,
                "longitude": lon,
                "is_public": to_is_public(row[cfg["pubpriv"]] if cfg["pubpriv"] < len(row) else None),
                "condition": condition,
                "notes": build_notes(row, cfg),
                "_priority": priority,
                "_sheet": sheet_name,
                "_excel_row": excel_row,
                "_adopter": (
                    clean(row[cfg["adopter_name"]]) if cfg["adopter_name"] < len(row) else "",
                    clean(row[cfg["adopter_phone"]]) if cfg["adopter_phone"] < len(row) else "",
                    clean(row[cfg["adopter_email"]]) if cfg["adopter_email"] < len(row) else "",
                ),
                "_skip_entry": skip_entry,
            }
            valid.append(record)
            n_valid += 1

        sheet_counts[sheet_name] = {"valid": n_valid, "skipped": n_skip, "status": status}

    return valid, skipped, warns, sheet_counts


def dedupe(valid, skipped):
    """Keep one record per ecoslo_num by priority; record dropped dups as skipped."""
    winners = {}
    for rec in valid:
        n = rec["ecoslo_num"]
        cur = winners.get(n)
        if cur is None:
            winners[n] = rec
        elif rec["_priority"] > cur["_priority"]:
            loser = cur
            loser["_skip_entry"]["reason"] = (
                f"duplicate ecoslo_num {n}: superseded by higher-priority sheet "
                f"'{rec['_sheet']}'"
            )
            skipped.append(loser["_skip_entry"])
            winners[n] = rec
        else:
            rec["_skip_entry"]["reason"] = (
                f"duplicate ecoslo_num {n}: kept row from '{cur['_sheet']}'"
            )
            skipped.append(rec["_skip_entry"])
    return winners


def build_members(winners):
    members = {}  # key -> member dict
    for rec in sorted(winners.values(), key=lambda r: r["ecoslo_num"]):
        raw_name, raw_phone, raw_email = rec["_adopter"]
        name = norm_ws(raw_name)
        phone = norm_ws(raw_phone)
        email = extract_email(raw_email)
        has_email = bool(email)
        has_name = bool(name) and name.lower() not in JUNK
        if not has_email and not has_name:
            continue
        key = email if has_email else "name::" + name.lower()
        m = members.get(key)
        if m is None:
            m = {
                "firstname": "", "lastname": "", "email": email if has_email else "",
                "phone": "", "real_email": has_email, "name_src": "", "trees": set(),
            }
            members[key] = m
        m["trees"].add(rec["ecoslo_num"])
        if not m["name_src"] and has_name:
            m["name_src"] = name
            parts = name.split()
            m["firstname"] = parts[0]
            m["lastname"] = " ".join(parts[1:]) if len(parts) > 1 else ""
        if not m["phone"] and phone and phone.lower() not in JUNK:
            m["phone"] = phone

    ordered = sorted(members.values(), key=lambda m: min(m["trees"]))
    seq = 0
    for m in ordered:
        if not m["real_email"]:
            seq += 1
            m["email"] = f"noemail+{seq}@ecoslo.invalid"
        # Flag rows a human should eyeball: synthesized email, no name at all, or
        # a "name" that is really a sentence/org (>4 words) the naive first-word/
        # remainder split will have mangled into firstname/lastname.
        wordy = len(m["name_src"].split()) > 4
        m["needs_review"] = (not m["real_email"]) or (not m["name_src"]) or wordy
    return ordered


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------
TREE_COLUMNS = [
    "ecoslo_num", "status", "species_name", "common_name", "funder",
    "date_planted", "address", "latitude", "longitude", "is_public",
    "condition", "notes",
]
MEMBER_COLUMNS = [
    "firstname", "lastname", "email", "phone", "role",
    "trees_assigned", "trees_count", "needs_review",
]
SKIPPED_COLUMNS = [
    "source_sheet", "excel_row", "reason", "raw_tree_num", "species",
    "common_name", "funder", "date_planted", "address", "raw_lat", "raw_long",
    "status_intended",
]


def write_trees_csv(path, winners):
    rows = sorted(winners.values(), key=lambda r: r["ecoslo_num"])
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=TREE_COLUMNS, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            out = {k: r[k] for k in TREE_COLUMNS}
            out["is_public"] = "true" if r["is_public"] else "false"
            w.writerow(out)
    return len(rows)


def write_members_csv(path, members):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=MEMBER_COLUMNS)
        w.writeheader()
        for m in members:
            trees = sorted(m["trees"])
            w.writerow({
                "firstname": m["firstname"],
                "lastname": m["lastname"],
                "email": m["email"],
                "phone": m["phone"],
                "role": "Tree Keeper",
                "trees_assigned": "{" + ",".join(str(t) for t in trees) + "}",
                "trees_count": len(trees),
                "needs_review": "true" if m["needs_review"] else "false",
            })
    return len(members)


def write_skipped_csv(path, skipped):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=SKIPPED_COLUMNS, extrasaction="ignore")
        w.writeheader()
        for s in skipped:
            w.writerow(s)
    return len(skipped)


def main():
    ap = argparse.ArgumentParser(description="Convert the Tree Planting Log workbook to Supabase CSVs.")
    ap.add_argument("input", help="Path to the .xlsx workbook")
    ap.add_argument("-o", "--outdir", default=".", help="Output directory (default: current dir)")
    args = ap.parse_args()

    if not os.path.isfile(args.input):
        sys.exit(f"Input file not found: {args.input}")
    os.makedirs(args.outdir, exist_ok=True)

    valid, skipped, warns, sheet_counts = parse_workbook(args.input)
    winners = dedupe(valid, skipped)
    members = build_members(winners)

    trees_path = os.path.join(args.outdir, "trees.csv")
    members_path = os.path.join(args.outdir, "members.csv")
    skipped_path = os.path.join(args.outdir, "skipped_rows.csv")

    n_trees = write_trees_csv(trees_path, winners)
    n_members = write_members_csv(members_path, members)
    n_skipped = write_skipped_csv(skipped_path, skipped)

    # ---- summary ----
    print("=" * 70)
    print("Tree Planting Log import")
    print("=" * 70)
    for name, c in sheet_counts.items():
        print(f"  {name:24} -> {c['status']:9}  valid:{c['valid']:4}  skipped:{c['skipped']:4}")
    placeholders = sum(1 for m in members if not m["real_email"])
    print("-" * 70)
    print(f"  trees.csv        : {n_trees} rows")
    print(f"  members.csv      : {n_members} rows  ({placeholders} with placeholder email / needs_review)")
    print(f"  skipped_rows.csv : {n_skipped} rows")
    if warns:
        print(f"  warnings         : {len(warns)}")
        for wmsg in warns[:10]:
            print(f"      - {wmsg}")
    print("-" * 70)
    print(f"  wrote {trees_path}")
    print(f"  wrote {members_path}")
    print(f"  wrote {skipped_path}")


if __name__ == "__main__":
    main()
